"""
数据处理模块：导入、检查、导出
"""
import io
import json
import math
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.calculator import CalculationResult, SurveyPoint, chain_segments, extract_zero_contour_segments
from utils.dxf_io import read_text_with_encodings
from utils.fileio import atomic_write_text
from version import __version__

def report_labels(result: CalculationResult) -> Dict[str, str]:
    """报告用语：与设计面比较，或两期对比。"""
    if result.is_compare:
        return {
            "natural": "前期高程",
            "design": "后期高程",
            "design_item": "比较面",
            "cut": "前期高于后期部分（挖除）",
            "fill": "后期高于前期部分（填筑）",
            "method": "两期TIN叠加（前期、后期三角网精确求交，边界精确裁剪）",
            "covered": "两期测量共同覆盖",
        }
    return {
        "natural": "实测高程",
        "design": "设计高程",
        "design_item": "设计高程",
        "cut": "实测高程高于设计高程部分",
        "fill": "实测高程低于设计高程部分",
        "method": "TIN三角网法（Delaunay剖分，边界精确裁剪）",
        "covered": "有测点覆盖",
    }


def _percent(value: Optional[float]) -> str:
    return "—" if value is None else f"{value:+.2f}%"


ISSUE_NAMES = {
    'missing_values': '缺失值',
    'duplicate_coords': '重复坐标',
    'duplicate_ids': '重复点号',
    'coord_outliers': '坐标异常',
    'elevation_outliers': '高程异常',
    'format_errors': '格式错误',
}


def _finite(value) -> float:
    """转为有限浮点数；空值、NaN、inf 视为无法解析。"""
    number = float(value)
    if not math.isfinite(number):
        raise ValueError("非有限数值")
    return number


class DataValidator:
    """数据验证器"""
    
    @staticmethod
    def validate_points(points: List[SurveyPoint]) -> Dict[str, Any]:
        """验证测量点数据"""
        issues = {
            'missing_values': [],
            'duplicate_coords': [],
            'duplicate_ids': [],
            'coord_outliers': [],
            'elevation_outliers': [],
            'format_errors': []
        }
        
        if not points:
            issues['format_errors'].append("没有数据")
            return issues
            
        coords_seen = {}
        ids_seen = {}
        elevations = [p.z for p in points]
        xs = [p.x for p in points]
        ys = [p.y for p in points]
        
        # 计算统计量用于异常值检测
        elev_mean, elev_std = np.mean(elevations), np.std(elevations)
        x_mean, x_std = np.mean(xs), np.std(xs)
        y_mean, y_std = np.mean(ys), np.std(ys)
        
        for i, pt in enumerate(points):
            # 缺失值检查
            if pd.isna(pt.x) or pd.isna(pt.y) or pd.isna(pt.z):
                issues['missing_values'].append(f"第{i+1}行: 坐标或高程缺失")
                
            # 重复坐标检查
            coord_key = (round(pt.x, 3), round(pt.y, 3))
            if coord_key in coords_seen:
                issues['duplicate_coords'].append(
                    f"第{i+1}行与第{coords_seen[coord_key]+1}行坐标重复: ({pt.x}, {pt.y})"
                )
            else:
                coords_seen[coord_key] = i
                
            # 重复点号检查
            if pt.id in ids_seen:
                issues['duplicate_ids'].append(f"第{i+1}行点号重复: {pt.id}")
            else:
                ids_seen[pt.id] = i
                
            # 坐标异常值检查 (3σ原则)
            if abs(pt.x - x_mean) > 3 * x_std and x_std > 0:
                issues['coord_outliers'].append(f"第{i+1}行 X坐标异常: {pt.x}")
            if abs(pt.y - y_mean) > 3 * y_std and y_std > 0:
                issues['coord_outliers'].append(f"第{i+1}行 Y坐标异常: {pt.y}")
                
            # 高程异常值检查
            if abs(pt.z - elev_mean) > 3 * elev_std and elev_std > 0:
                issues['elevation_outliers'].append(f"第{i+1}行 高程异常: {pt.z}")
                
        return issues
    
    @staticmethod
    def auto_detect_columns(df: pd.DataFrame) -> Dict[str, str]:
        """自动识别列映射"""
        cols = {c.lower().strip(): c for c in df.columns}
        mapping = {}
        
        # 点号列
        for key in ['点号', 'id', 'pid', 'point', '编号', 'name']:
            if key in cols:
                mapping['id'] = cols[key]
                break
                
        # X坐标列
        for key in ['x', 'x坐标', '横坐标', '东向', 'easting', '经度', 'lon', 'longitude']:
            if key in cols:
                mapping['x'] = cols[key]
                break
                
        # Y坐标列
        for key in ['y', 'y坐标', '纵坐标', '北向', 'northing', '纬度', 'lat', 'latitude']:
            if key in cols:
                mapping['y'] = cols[key]
                break
                
        # 高程列
        for key in ['z', '实测高程', '高程', 'elevation', 'alt', 'altitude', 'h']:
            if key in cols:
                mapping['z'] = cols[key]
                break

        for key in ['设计高程', '设计标高', 'design_z', 'designz', 'designelev', 'design_elev']:
            if key in cols:
                mapping['design_z'] = cols[key]
                break
                
        return mapping


class DataImporter:
    """数据导入器"""

    @staticmethod
    def complete_column_mapping(df: pd.DataFrame, mapping: Optional[Dict[str, str]] = None) -> Dict[str, str]:
        """补齐缺失的 id/x/y/z 映射：按顺序取尚未使用的列，不覆盖已识别出的列。

        列数不够时优先保证 X/Y/Z，点号缺失会在解析时自动编号。
        """
        completed = dict(mapping or {})
        used = set(completed.values())
        unused = [column for column in df.columns if column not in used]
        required = [key for key in ("x", "y", "z") if key not in completed]
        if "id" not in completed and len(unused) > len(required):
            completed["id"] = unused.pop(0)
        for key in required:
            if unused:
                completed[key] = unused.pop(0)
        return completed

    @staticmethod
    def points_from_dataframe(
        df: pd.DataFrame, mapping: Dict[str, str]
    ) -> Tuple[List[SurveyPoint], List[str]]:
        """按列映射解析测点；无法解析的行记入 format_errors，不静默丢弃。"""
        points: List[SurveyPoint] = []
        format_errors: List[str] = []
        id_column = mapping.get("id")
        x_column = mapping.get("x")
        y_column = mapping.get("y")
        z_column = mapping.get("z")
        design_column = mapping.get("design_z")

        # 点号直接从列里取：逐行取时整数会随其他列升成浮点，点号 7 就成了 “7.0”
        id_values = None
        if id_column is not None:
            id_series = df[id_column]
            if isinstance(id_series, pd.DataFrame):   # 表头里有重名的列，取第一列
                id_series = id_series.iloc[:, 0]
            id_values = id_series.tolist()

        for row_number, (index, row) in enumerate(df.iterrows(), start=1):
            try:
                if x_column is None or y_column is None or z_column is None:
                    raise KeyError("缺少 X/Y/Z 列映射")
                raw_id = id_values[row_number - 1] if id_values is not None else None
                if raw_id is None or pd.isna(raw_id) or str(raw_id).strip() == "":
                    point_id = f"P{row_number}"
                elif isinstance(raw_id, float) and raw_id.is_integer():
                    point_id = str(int(raw_id))   # 列里有空点号时整列是浮点
                else:
                    point_id = str(raw_id).strip()
                design_z = 0.0
                has_design_z = False
                if design_column is not None:
                    raw_design = row[design_column]
                    if pd.notna(raw_design) and str(raw_design).strip() != "":
                        design_z = _finite(raw_design)
                        has_design_z = True
                z_value = _finite(row[z_column])
                pt = SurveyPoint(
                    id=point_id,
                    x=_finite(row[x_column]),
                    y=_finite(row[y_column]),
                    z=z_value,
                    design_z=design_z,
                    delta_z=(z_value - design_z) if has_design_z else 0.0,
                    has_design_z=has_design_z,
                )
                points.append(pt)
            except (ValueError, TypeError, KeyError):
                format_errors.append(f"第{row_number}行: 坐标或高程为空或无法解析，已跳过")
        return points, format_errors
    
    @staticmethod
    def import_file(filepath: str) -> Tuple[List[SurveyPoint], Dict[str, Any], pd.DataFrame]:
        """导入文件，返回(点列表, 验证结果, 原始DataFrame)"""
        path = Path(filepath)
        suffix = path.suffix.lower()
        
        # 没有表头的 CSV 按“点号,X,Y,高程”文本读；交给 pandas 会把第一个测点当成表头丢掉
        as_text = suffix in ['.txt', '.dat'] or (suffix == '.csv' and not DataImporter._csv_has_header(filepath))

        try:
            if suffix in ['.xlsx', '.xls']:
                df = pd.read_excel(filepath)
            elif suffix == '.csv' and not as_text:
                # 尝试多种编码
                for enc in ['utf-8', 'gbk', 'gb2312', 'utf-16']:
                    try:
                        df = pd.read_csv(filepath, encoding=enc)
                        break
                    except UnicodeDecodeError:
                        continue
                else:
                    raise ValueError("无法识别CSV编码")
            elif as_text:
                points, format_errors, df = DataImporter.import_survey_text(
                    filepath, cass=(suffix == '.dat')
                )
                issues = DataValidator.validate_points(points)
                issues["format_errors"].extend(format_errors)
                return points, issues, df
            else:
                raise ValueError(f"不支持的文件格式: {suffix}")
                
        except Exception as e:
            raise ValueError(f"文件读取失败: {e}")
            
        if df.empty:
            raise ValueError("文件为空")
            
        mapping = DataImporter.complete_column_mapping(
            df, DataValidator.auto_detect_columns(df)
        )
        points, format_errors = DataImporter.points_from_dataframe(df, mapping)
        issues = DataValidator.validate_points(points)
        issues["format_errors"].extend(format_errors)
        return points, issues, df

    @staticmethod
    def _csv_has_header(filepath: str) -> bool:
        """CSV 的第一行是不是表头（而不是第一个测点）。读不了的文件当作有表头，由后面的读取报错。"""
        try:
            text = read_text_with_encodings(filepath)
        except OSError:
            return True
        for raw in text.splitlines():
            line = raw.strip()
            if line:
                return DataImporter._is_survey_header(DataImporter._split_survey_fields(line))
        return True

    @staticmethod
    def _split_survey_fields(line: str) -> List[str]:
        if "," in line:
            return [part.strip() for part in line.split(",")]
        return line.split()

    @staticmethod
    def _is_survey_header(fields: List[str]) -> bool:
        joined = "".join(fields).lower()
        if any(keyword in joined for keyword in ("点号", "点名", "编号", "实测", "高程", "坐标", "northing", "easting")):
            return True
        if len(fields) >= 5:
            try:
                float(fields[2])
                float(fields[3])
                float(fields[4])
                return False
            except ValueError:
                return True
        if len(fields) >= 4:
            try:
                float(fields[1])
                float(fields[2])
                float(fields[-1])
                return False
            except ValueError:
                return True
        if len(fields) == 3:
            try:
                float(fields[0])
                float(fields[1])
                float(fields[2])
                return False
            except ValueError:
                return True
        return False

    @staticmethod
    def _looks_like_cass(fields: List[str]) -> bool:
        if len(fields) < 5:
            return False
        try:
            float(fields[2])
            float(fields[3])
            float(fields[4])
        except (TypeError, ValueError):
            return False
        try:
            float(fields[1])
            return False
        except (TypeError, ValueError):
            return True

    @staticmethod
    def _points_to_dataframe(points: List[SurveyPoint]) -> pd.DataFrame:
        rows = {
            "点号": [point.id for point in points],
            "X坐标": [point.x for point in points],
            "Y坐标": [point.y for point in points],
            "实测高程": [point.z for point in points],
        }
        if any(point.has_design_z for point in points):
            rows["设计高程"] = [point.design_z for point in points]
        return pd.DataFrame(rows)

    @staticmethod
    def import_survey_text(
        filepath: str, cass: bool = False
    ) -> Tuple[List[SurveyPoint], List[str], pd.DataFrame]:
        """导入 CASS .dat 或点号/X/Y/Z 的空格、逗号分隔文本。"""
        text = read_text_with_encodings(filepath)
        lines = []
        for raw in text.splitlines():
            line = raw.strip()
            if not line or line.startswith(("#", "*", "//")):
                continue
            if line.upper() in {"BEGIN", "END"}:
                continue
            lines.append(line)
        if not lines:
            raise ValueError("文件为空")

        first_fields = DataImporter._split_survey_fields(lines[0])
        header = DataImporter._is_survey_header(first_fields)
        body = lines[1:] if header else lines
        if not body:
            raise ValueError("文件为空")
        sample = DataImporter._split_survey_fields(body[0])
        use_cass = cass or DataImporter._looks_like_cass(sample)

        if header and not use_cass:
            rows = [DataImporter._split_survey_fields(line) for line in body]
            width = max(len(first_fields), max((len(row) for row in rows), default=0))
            columns = list(first_fields) + [f"col{i}" for i in range(len(first_fields), width)]
            padded = [row + [""] * (width - len(row)) for row in rows]
            df = pd.DataFrame(padded, columns=columns[:width])
            mapping = DataImporter.complete_column_mapping(df, DataValidator.auto_detect_columns(df))
            points, format_errors = DataImporter.points_from_dataframe(df, mapping)
            return points, format_errors, df

        points: List[SurveyPoint] = []
        format_errors: List[str] = []
        for row_number, line in enumerate(body, start=1):
            fields = DataImporter._split_survey_fields(line)
            try:
                if use_cass:
                    if len(fields) >= 5:
                        point_id, easting, northing, height = fields[0], fields[2], fields[3], fields[4]
                    elif len(fields) == 4:
                        point_id, easting, northing, height = fields[0], fields[1], fields[2], fields[3]
                    elif len(fields) == 3:
                        point_id, easting, northing, height = f"P{row_number}", fields[0], fields[1], fields[2]
                    else:
                        raise ValueError("字段不足")
                    point = SurveyPoint(
                        id=str(point_id),
                        x=_finite(easting),
                        y=_finite(northing),
                        z=_finite(height),
                    )
                else:
                    if len(fields) >= 4:
                        point = SurveyPoint(
                            id=str(fields[0]),
                            x=_finite(fields[1]),
                            y=_finite(fields[2]),
                            z=_finite(fields[3]),
                        )
                    elif len(fields) == 3:
                        point = SurveyPoint(
                            id=f"P{row_number}",
                            x=_finite(fields[0]),
                            y=_finite(fields[1]),
                            z=_finite(fields[2]),
                        )
                    else:
                        raise ValueError("字段不足")
                points.append(point)
            except (ValueError, TypeError):
                format_errors.append(f"第{row_number}行: 坐标或高程为空或无法解析，已跳过")
        if not points:
            raise ValueError("没有可解析的测量点")
        return points, format_errors, DataImporter._points_to_dataframe(points)


class DataExporter:
    """数据导出器。

    导出失败时直接抛出异常，由界面把原因展示给用户。
    """

    DETAIL_HEADERS = [
        '三角形编号', '顶点1点号', '顶点1X', '顶点1Y', '顶点1实测高程', '顶点1设计高程', '顶点1高差',
        '顶点2点号', '顶点2X', '顶点2Y', '顶点2实测高程', '顶点2设计高程', '顶点2高差',
        '顶点3点号', '顶点3X', '顶点3Y', '顶点3实测高程', '顶点3设计高程', '顶点3高差',
        '边界内面积(m²)', '平均高差(m)', '挖方量(m³)', '填方量(m³)', '净体积(m³)', '是否混合', '是否被边界裁剪',
    ]
    GRID_CORNER_HEADERS = ['角点(行-列)', '行', '列', 'X', 'Y', '{natural}', '{design}', '施工高度(m)', '备注']
    GRID_CELL_HEADERS = ['方格(行-列)', '左下角点', '右下角点', '右上角点', '左上角点',
                         '计算面积(m²)', '挖方量(m³)', '填方量(m³)', '净方量(m³)', '是否被边界裁剪']

    def __init__(self, project_name: str = "土方计算项目"):
        self.project_name = project_name
        self.calc_date = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    @classmethod
    def detail_headers(cls, result: CalculationResult) -> List[str]:
        labels = report_labels(result)
        return [
            header.replace("实测高程", labels["natural"]).replace("设计高程", labels["design"])
            for header in cls.DETAIL_HEADERS
        ]

    @staticmethod
    def grid_corner_rows(result: CalculationResult) -> List[List[Any]]:
        grid = result.grid_check
        if grid is None:
            return []
        rows = []
        for corner in grid.corners:
            if not corner.valid:
                rows.append([corner.name, corner.row, corner.col, corner.x, corner.y, '', '', '', '无法取值'])
                continue
            note = '测区外，按相邻三角形外推，仅用于方格内插值' if corner.extrapolated else ''
            rows.append([corner.name, corner.row, corner.col, corner.x, corner.y,
                         round(corner.natural, 3), round(corner.design, 3), round(corner.height, 3), note])
        return rows

    @staticmethod
    def grid_cell_rows(result: CalculationResult) -> List[List[Any]]:
        grid = result.grid_check
        if grid is None:
            return []
        return [
            [cell.name, *cell.corners, round(cell.area, 3), round(cell.cut, 3), round(cell.fill, 3),
             round(cell.net, 3), '是' if cell.clipped else '否']
            for cell in grid.cells
        ]

    @staticmethod
    def _detail_rows(result: CalculationResult):
        for tri in result.triangles:
            if tri.is_boundary:
                continue
            row = [tri.id]
            for point in tri.vertex_points:
                row += [point.id, point.x, point.y, point.z, point.design_z, round(point.delta_z, 4)]
            row += [
                round(tri.area, 3),
                '' if tri.is_mixed else round(tri.avg_delta_z, 3),
                round(tri.cut_volume, 3),
                round(tri.fill_volume, 3),
                round(tri.volume, 3),
                '是' if tri.is_mixed else '否',
                '是' if tri.is_clipped else '否',
            ]
            yield row

    @staticmethod
    def summary_rows(result: CalculationResult) -> List[List[Any]]:
        """汇总表各行：项目、数值、单位、说明。界面汇总页与报告共用。"""
        labels = report_labels(result)
        rows = [
            ['总挖方量', round(result.total_cut, 3), 'm³', labels["cut"]],
            ['总填方量', round(result.total_fill, 3), 'm³', labels["fill"]],
            ['挖填差值', round(result.net_volume, 3), 'm³', '正值为挖方大（需外运），负值为填方大（需借方）'],
            ['计算面积', round(result.computed_area, 3), 'm²', f'计算边界内且{labels["covered"]}的水平面积'],
        ]
        if result.boundary_area > 0:
            rows.append(['计算边界面积', round(result.boundary_area, 3), 'm²', ''])
            rows.append(['测点覆盖率', round(result.coverage_ratio * 100, 2), '%', '低于 99.5% 时请核对边界'])
        if result.design_description:
            # 斜面、两期对比的说明较长：数值列写类型，完整说明放在说明列
            short = "两期对比" if result.is_compare else result.design_description.split("：", 1)[0]
            rows.append([labels["design_item"], short, '', result.design_description])
        else:
            rows.append([labels["design_item"], result.design_text, '', ''])
        if result.is_compare:
            rows.append(['后期测点数', result.compare_point_count, '个', ''])
        rows += [
            ['三角形数', result.triangle_count, '个',
             '两期三角网叠加后的计算单元' if result.is_compare else '参与计算的三角形'],
            ['混合三角形', result.mixed_triangle_count, '个', '跨越零填挖线，已按零线分割'],
            ['边界裁剪三角形', result.clipped_triangle_count, '个', '跨越计算边界，只计边界内部分'],
        ]
        grid = result.grid_check
        if grid is not None:
            comparison = {label: (diff, relative) for label, _, _, diff, relative in grid.comparison(result)}
            rows += [
                ['方格网校核', f'边长 {grid.spacing:g} m', '', f'计入方格 {len(grid.cells)} 个，三角棱柱体法'],
                ['方格网法挖方', round(grid.total_cut, 3), 'm³', f'与 TIN 法相差 {_percent(comparison["挖方"][1])}'],
                ['方格网法填方', round(grid.total_fill, 3), 'm³', f'与 TIN 法相差 {_percent(comparison["填方"][1])}'],
                ['方格网法挖填差值', round(grid.net_volume, 3), 'm³',
                 f'与 TIN 法相差 {comparison["挖填差值"][0]:+.3f} m³'],
            ]
        return rows

    @staticmethod
    def check_rows(points: Optional[List[SurveyPoint]], result: CalculationResult, limit: int = 30) -> List[List[Any]]:
        """异常数据检查页：导入检查结果 + 计算过程中的提示。"""
        rows: List[List[Any]] = []
        if points is not None:
            issues = DataValidator.validate_points(points)
            for key, name in ISSUE_NAMES.items():
                items = issues.get(key, [])
                detail = "；".join(items[:limit]) + (f"；……共 {len(items)} 条" if len(items) > limit else "")
                rows.append([name, len(items), detail or "未发现"])
        rows.append(["计算提示", len(result.warnings), "\n".join(result.warnings) or "无"])
        if result.grid_check is not None:
            grid_warnings = result.grid_check.warnings
            rows.append(["方格网校核提示", len(grid_warnings), "\n".join(grid_warnings) or "无"])
        return rows

    def export_summary_excel(self, result: CalculationResult, filepath: str,
                             points: Optional[List[SurveyPoint]] = None) -> bool:
        """导出汇总表、三角形明细和异常数据检查。

        使用只写模式并只给表头设置样式，大数据量时速度约为逐格设置样式的 3 倍。
        """
        wb = openpyxl.Workbook(write_only=True)
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center = Alignment(horizontal='center', vertical='center')

        def header(ws, values):
            cells = []
            for value in values:
                cell = WriteOnlyCell(ws, value=value)
                cell.font, cell.fill, cell.alignment = header_font, header_fill, center
                cells.append(cell)
            ws.append(cells)

        def bold(ws, value, size=11):
            cell = WriteOnlyCell(ws, value=value)
            cell.font = Font(bold=True, size=size)
            return cell

        labels = report_labels(result)
        ws1 = wb.create_sheet("土方量汇总表")
        for column, width in zip("ABCDEF", (18, 26, 10, 44, 18, 18)):
            ws1.column_dimensions[column].width = width
        ws1.append([bold(ws1, self.project_name, 14)])
        ws1.append([bold(ws1, "土方量计算汇总表")])
        ws1.append([])
        ws1.append(["计算日期", self.calc_date, "", "计算方法", labels["method"]])
        ws1.append([labels["design_item"], result.design_text, "", "计算边界顶点数", len(result.boundary_points or [])])
        ws1.append([])
        header(ws1, ['项目', '数值', '单位', '说明'])
        for row in self.summary_rows(result):
            ws1.append(row)

        ws2 = wb.create_sheet("三角形计算明细")
        ws2.freeze_panes = "B2"
        detail_headers = self.detail_headers(result)
        for index in range(len(detail_headers)):
            ws2.column_dimensions[get_column_letter(index + 1)].width = 14
        header(ws2, detail_headers)
        for row in self._detail_rows(result):
            ws2.append(row)

        ws3 = wb.create_sheet("异常数据检查")
        ws3.column_dimensions["A"].width = 14
        ws3.column_dimensions["B"].width = 10
        ws3.column_dimensions["C"].width = 100
        header(ws3, ["检查项目", "发现数量", "详细信息"])
        for row in self.check_rows(points, result):
            ws3.append(row)

        grid = result.grid_check
        if grid is not None:
            ws4 = wb.create_sheet("方格网校核")
            for column, width in zip("ABCDEF", (16, 16, 16, 16, 12, 30)):
                ws4.column_dimensions[column].width = width
            ws4.append([bold(ws4, f"方格网法校核（边长 {grid.spacing:g} m，三角棱柱体法）", 12)])
            ws4.append(["方格网原点", f"({grid.origin[0]:.3f}, {grid.origin[1]:.3f})", "行 × 列",
                        f"{grid.rows} × {grid.cols}", "计入方格", len(grid.cells)])
            ws4.append([])
            header(ws4, ["项目", "TIN 法", "方格网法", "差值", "相对差"])
            for label, tin_value, grid_value, diff, relative in grid.comparison(result):
                ws4.append([label, round(tin_value, 3), round(grid_value, 3), round(diff, 3), _percent(relative)])
            ws4.append([])
            ws4.append(["说明", "角点高程由三角网线性插值，每个方格沿左下—右上对角线分成两个三角形，"
                        "按边界精确裁剪、按零线分割。与 TIN 法的差异反映方格取样密度的影响。"])
            for warning in grid.warnings:
                ws4.append(["提示", warning])

            ws5 = wb.create_sheet("方格网角点")
            ws5.freeze_panes = "A2"
            header(ws5, [h.format(natural=labels["natural"], design=labels["design"]) for h in self.GRID_CORNER_HEADERS])
            for row in self.grid_corner_rows(result):
                ws5.append(row)

            ws6 = wb.create_sheet("方格网方格")
            ws6.freeze_panes = "A2"
            header(ws6, self.GRID_CELL_HEADERS)
            for row in self.grid_cell_rows(result):
                ws6.append(row)

        # 先完整写入内存再落盘：目标文件被占用时不会留下半成品和临时文件
        buffer = io.BytesIO()
        wb.save(buffer)
        Path(filepath).write_bytes(buffer.getvalue())
        return True

    def export_triangles_csv(self, result: CalculationResult, filepath: str) -> bool:
        """导出三角形明细CSV"""
        df = pd.DataFrame(list(self._detail_rows(result)), columns=self.detail_headers(result))
        df.to_csv(filepath, index=False, encoding='utf-8-sig')
        return True

    def export_boundary_dxf(self, result: CalculationResult, filepath: str) -> bool:
        """导出计算边界（闭合多段线）和零填挖线到 DXF。

        采用 R12 格式的 POLYLINE，AutoCAD、CASS 等都能直接打开。
        """
        parts = ["0\nSECTION\n2\nHEADER\n9\n$ACADVER\n1\nAC1009\n0\nENDSEC\n",
                 "0\nSECTION\n2\nENTITIES\n"]
        if result.boundary_points:
            parts.append(self._dxf_polyline("BOUNDARY", result.boundary_points, closed=True))
        for polyline in chain_segments(extract_zero_contour_segments(result)):
            if len(polyline) >= 2:
                parts.append(self._dxf_polyline("ZERO_CONTOUR", polyline, closed=False))
        if result.grid_check is not None:
            parts.extend(self._dxf_grid(result.grid_check))
        parts.append("0\nENDSEC\n0\nEOF\n")
        with open(filepath, "w", encoding="utf-8") as handle:
            handle.write("".join(parts))
        return True

    @staticmethod
    def _dxf_polyline(layer: str, points, closed: bool) -> str:
        lines = [f"0\nPOLYLINE\n8\n{layer}\n66\n1\n10\n0.0\n20\n0.0\n30\n0.0\n70\n{1 if closed else 0}\n"]
        for x, y in points:
            lines.append(f"0\nVERTEX\n8\n{layer}\n10\n{float(x)!r}\n20\n{float(y)!r}\n30\n0.0\n")
        lines.append(f"0\nSEQEND\n8\n{layer}\n")
        return "".join(lines)

    @staticmethod
    def _dxf_text(layer: str, x: float, y: float, height: float, text: str) -> str:
        return f"0\nTEXT\n8\n{layer}\n10\n{float(x)!r}\n20\n{float(y)!r}\n30\n0.0\n40\n{float(height)!r}\n1\n{text}\n"

    @classmethod
    def _dxf_grid(cls, grid) -> List[str]:
        """方格网图：GRID 方格线、GRID_HEIGHT 角点施工高度、GRID_ELEV 角点 地面/设计 高程、GRID_VOLUME 方格方量。

        方格线铺满计算范围的外接矩形；方量只计计算范围以内的部分。

        文字只含数字和符号，R12 下无需处理中文编码。
        """
        a = grid.spacing
        x0, y0 = grid.origin
        x1, y1 = x0 + grid.cols * a, y0 + grid.rows * a
        parts = [cls._dxf_polyline("GRID", [(x0 + i * a, y0), (x0 + i * a, y1)], closed=False)
                 for i in range(grid.cols + 1)]
        parts += [cls._dxf_polyline("GRID", [(x0, y0 + j * a), (x1, y0 + j * a)], closed=False)
                  for j in range(grid.rows + 1)]
        size = a * 0.08
        for corner in grid.corners:
            # 测区外外推的角点不是实测值，图上不标注
            if not corner.valid or corner.extrapolated:
                continue
            parts.append(cls._dxf_text("GRID_HEIGHT", corner.x + size * 0.5, corner.y + size * 0.5, size,
                                       f"{corner.height:+.2f}"))
            parts.append(cls._dxf_text("GRID_ELEV", corner.x + size * 0.5, corner.y - size * 1.5, size * 0.8,
                                       f"{corner.natural:.2f}/{corner.design:.2f}"))
        for cell in grid.cells:
            cx = x0 + (cell.col - 0.5) * a
            cy = y0 + (cell.row - 0.5) * a
            if cell.cut > 0 and cell.fill > 0:
                text = f"+{cell.cut:.1f}/-{cell.fill:.1f}"
            else:
                text = f"{cell.net:+.1f}"
            parts.append(cls._dxf_text("GRID_VOLUME", cx - a * 0.3, cy, size, text))
        return parts

    def generate_report_text(self, result: CalculationResult, point_count: int) -> str:
        """生成文本报告"""
        labels = report_labels(result)
        if result.net_volume > 0:
            balance = "(挖方大于填方，需外运)"
        elif result.net_volume < 0:
            balance = "(填方大于挖方，需借方)"
        else:
            balance = "(挖填平衡)"
        area_lines = f"计算面积: {result.computed_area:.3f} m²"
        if result.boundary_area > 0:
            area_lines += (
                f"\n计算边界面积: {result.boundary_area:.3f} m²"
                f"（测点覆盖率 {result.coverage_ratio * 100:.2f}%）"
            )
        warnings = "\n".join(f"- {item}" for item in result.warnings) or "无"
        point_lines = f"测量点数: {point_count}"
        if result.is_compare:
            point_lines = f"前期测点数: {point_count}\n后期测点数: {result.compare_point_count}"
            notes = (
                "1. 前期、后期测点分别构建TIN三角网，两网精确叠加后，每个计算单元内两期地面均为平面。\n"
                "2. 高差 = 前期高程 − 后期高程；正为挖除，负为填筑。\n"
                "3. 跨越计算边界的单元按边界精确裁剪；只计算两期都有测点覆盖的部分。\n"
                "4. 混合挖填单元已按零填挖线分割计算，避免正负高差抵消。\n"
                "5. 计算结果仅供工程参考，正式计量请以复核成果为准。"
            )
        else:
            notes = (
                "1. 本成果基于实测点构建TIN三角网，按水平投影面积×平均高差法计算。\n"
                "2. 跨越计算边界的三角形按边界精确裁剪，只计边界以内部分。\n"
                "3. 混合挖填三角形已按零填挖线分割计算，避免正负高差抵消。\n"
                "4. 所有三角形均保留面积、顶点高程、高差及分项体积，可逐项复核。\n"
                "5. 计算结果仅供工程参考，正式计量请以复核成果为准。"
            )
        grid_section = ""
        grid = result.grid_check
        if grid is not None:
            lines = [
                f"方格边长: {grid.spacing:g} m，方格网 {grid.rows} 行 × {grid.cols} 列，计入方格 {len(grid.cells)} 个",
                "计算方法: 角点高程由三角网插值，每格沿对角线分为两个三角形（三角棱柱体法）",
            ]
            for label, tin_value, grid_value, diff, relative in grid.comparison(result):
                unit = "m²" if label == "计算面积" else "m³"
                lines.append(
                    f"{label}: 方格网法 {grid_value:.3f} {unit}，TIN 法 {tin_value:.3f} {unit}，"
                    f"相差 {diff:+.3f} {unit}（{_percent(relative)}）"
                )
            lines += [f"- {item}" for item in grid.warnings]
            grid_section = "\n【方格网校核】\n" + "\n".join(lines) + "\n"
        return f"""
============================================================
              {self.project_name} - 土方计算报告
============================================================

【项目基本信息】
计算日期: {self.calc_date}
计算方法: {labels["method"]}
{point_lines}
三角形总数: {result.triangle_count}
混合挖填三角形: {result.mixed_triangle_count}
边界裁剪三角形: {result.clipped_triangle_count}
计算边界顶点数: {len(result.boundary_points) if result.boundary_points else 0}
{labels["design_item"]}: {result.design_text}
{area_lines}

【计算结果汇总】
总挖方量: {result.total_cut:.3f} m³
总填方量: {result.total_fill:.3f} m³
挖填差值: {result.net_volume:.3f} m³
{balance}
{grid_section}
【计算提示】
{warnings}

【说明】
{notes}

============================================================
"""

    def export_pdf_report(
        self,
        result: CalculationResult,
        point_count: int,
        points: List[SurveyPoint],
        boundary: List[Tuple[float, float]],
        filepath: str,
    ) -> bool:
        """导出含挖填汇总与成果图的 PDF 计算书。"""
        from matplotlib.backends.backend_pdf import PdfPages
        from matplotlib.figure import Figure
        from utils.plotter import create_standalone_figure, setup_chinese_font

        setup_chinese_font()
        with PdfPages(filepath) as pdf:
            summary = Figure(figsize=(8.27, 11.69))
            axis = summary.add_subplot(111)
            axis.axis("off")
            axis.text(
                0.06,
                0.97,
                self.generate_report_text(result, point_count).strip(),
                va="top",
                ha="left",
                fontsize=9,
                wrap=True,
                transform=axis.transAxes,
            )
            pdf.savefig(summary)

            figure = create_standalone_figure(result, points, boundary, self.project_name)
            pdf.savefig(figure)
        return True


PROJECT_FORMAT_VERSION = 2
DESIGN_MODES = ("flat", "plane", "compare")


def _point_rows(points: List[SurveyPoint]) -> List[Dict[str, Any]]:
    return [
        {
            "id": point.id,
            "x": point.x,
            "y": point.y,
            "z": point.z,
            "design_z": point.design_z,
            "has_design_z": bool(point.has_design_z),
        }
        for point in points
    ]


def _points_from_rows(rows) -> List[SurveyPoint]:
    points = []
    for row in rows or []:
        design_z = float(row.get("design_z", 0.0))
        has_design_z = bool(row.get("has_design_z", "design_z" in row))
        z_value = float(row["z"])
        points.append(
            SurveyPoint(
                id=str(row["id"]),
                x=float(row["x"]),
                y=float(row["y"]),
                z=z_value,
                design_z=design_z,
                delta_z=z_value - design_z if has_design_z else 0.0,
                has_design_z=has_design_z,
            )
        )
    return points


def save_project(
    filepath: str,
    points: List[SurveyPoint],
    boundary: List[Tuple[float, float]],
    design_elevation: float = 0.0,
    use_partition: bool = False,
    partition: Optional[Dict[str, float]] = None,
    project_name: str = "TIN土方计算项目",
    design_mode: str = "flat",
    design_plane: Optional[Dict[str, float]] = None,
    compare_points: Optional[List[SurveyPoint]] = None,
    compare_source: str = "",
    grid_enabled: bool = False,
    grid_spacing: Optional[float] = None,
) -> None:
    """保存可再次打开的工程文件（测点、边界、设计面、后期测点、方格网设置）。"""
    if design_mode not in DESIGN_MODES:
        raise ValueError(f"未知的设计面类型: {design_mode}")
    payload = {
        "format": "tin-earthwork-project",
        "format_version": PROJECT_FORMAT_VERSION,
        "app_version": __version__,
        "project_name": project_name,
        "design_mode": design_mode,
        "design_elevation": float(design_elevation),
        "design_plane": {key: float(value) for key, value in (design_plane or {}).items()} or None,
        "use_partition": bool(use_partition),
        "partition": {str(key): float(value) for key, value in (partition or {}).items()},
        "grid_check": {"enabled": bool(grid_enabled), "spacing": None if grid_spacing is None else float(grid_spacing)},
        "compare_source": compare_source,
        "boundary": [[float(x), float(y)] for x, y in boundary],
        "points": _point_rows(points),
        "compare_points": _point_rows(compare_points or []),
    }
    atomic_write_text(filepath, json.dumps(payload, ensure_ascii=False, indent=2))


def load_project(filepath: str) -> Dict[str, Any]:
    """读取工程文件（兼容 1.x 的格式），返回测点对象和计算设置。"""
    with open(filepath, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if payload.get("format") not in (None, "tin-earthwork-project"):
        raise ValueError("不是 TIN 土方工程文件")
    if int(payload.get("format_version") or 1) > PROJECT_FORMAT_VERSION:
        raise ValueError("工程文件由更新版本的程序保存，请升级后再打开")
    boundary = [(float(x), float(y)) for x, y in (payload.get("boundary") or [])]
    partition = {str(key): float(value) for key, value in (payload.get("partition") or {}).items()}
    compare_points = _points_from_rows(payload.get("compare_points"))
    design_mode = payload.get("design_mode") or ("compare" if compare_points else "flat")
    if design_mode not in DESIGN_MODES:
        design_mode = "flat"
    grid = payload.get("grid_check") or {}
    return {
        "project_name": payload.get("project_name") or "TIN土方计算项目",
        "design_mode": design_mode,
        "design_elevation": float(payload.get("design_elevation") or 0.0),
        "design_plane": payload.get("design_plane") or None,
        "use_partition": bool(payload.get("use_partition")),
        "partition": partition,
        "grid_enabled": bool(grid.get("enabled")),
        "grid_spacing": None if grid.get("spacing") is None else float(grid["spacing"]),
        "compare_source": payload.get("compare_source") or "",
        "boundary": boundary,
        "points": _points_from_rows(payload.get("points")),
        "compare_points": compare_points,
    }
