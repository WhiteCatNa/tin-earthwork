import pandas as pd
import openpyxl

from core.calculator import SurveyPoint, TINEarthworkCalculator
from utils.data_handler import DataExporter, DataImporter, DataValidator, load_project, save_project
from utils.dxf_io import import_boundary_from_dxf


def _result():
    points = [
        SurveyPoint("A", 0, 0, 12),
        SurveyPoint("B", 1, 0, 8),
        SurveyPoint("C", 0, 1, 10),
    ]
    calculator = TINEarthworkCalculator(10)
    calculator.add_points(points)
    return calculator.run_full_calculation(), points


def test_import_csv_and_detect_columns(tmp_path):
    source = tmp_path / "points.csv"
    pd.DataFrame({
        "id": ["A", "B", "C"],
        "x": [0, 1, 0],
        "y": [0, 0, 1],
        "z": [10, 11, 12],
    }).to_csv(source, index=False)

    points, issues, _ = DataImporter.import_file(source)
    assert [point.id for point in points] == ["A", "B", "C"]
    assert not any(issues.values())


def test_validator_reports_duplicate_ids_and_coordinates():
    issues = DataValidator.validate_points([
        SurveyPoint("A", 0, 0, 10),
        SurveyPoint("A", 0, 0, 11),
    ])
    assert issues["duplicate_ids"]
    assert issues["duplicate_coords"]


def test_import_reports_unparseable_rows(tmp_path):
    source = tmp_path / "bad.csv"
    source.write_text("id,x,y,z\nA,0,0,10\nB,bad,1,11\nC,0,1,12\n", encoding="utf-8")

    points, issues, _ = DataImporter.import_file(source)

    assert [point.id for point in points] == ["A", "C"]
    assert any("第2行" in item for item in issues["format_errors"])


def test_import_chinese_headers_out_of_order(tmp_path):
    source = tmp_path / "points.xlsx"
    pd.DataFrame({
        "实测高程": [10, 11, 12],
        "X坐标": [0, 1, 0],
        "Y坐标": [0, 0, 1],
        "点号": ["A", "B", "C"],
    }).to_excel(source, index=False)

    points, issues, _ = DataImporter.import_file(source)

    by_id = {point.id: point for point in points}
    assert by_id["A"].x == 0
    assert by_id["B"].x == 1
    assert by_id["C"].z == 12
    assert not issues["format_errors"]


def test_import_design_elevation_column(tmp_path):
    source = tmp_path / "design.csv"
    pd.DataFrame({
        "点号": ["A", "B", "C"],
        "X坐标": [0, 1, 0],
        "Y坐标": [0, 0, 1],
        "实测高程": [12, 8, 10],
        "设计高程": [11, 9, 10],
    }).to_csv(source, index=False)

    points, issues, _ = DataImporter.import_file(source)
    by_id = {point.id: point for point in points}

    assert by_id["A"].has_design_z
    assert by_id["A"].design_z == 11
    assert by_id["B"].design_z == 9
    assert by_id["C"].delta_z == 0
    assert not issues["format_errors"]


def test_project_file_roundtrip_restores_points_boundary_and_design(tmp_path):
    points = [
        SurveyPoint("A", 0, 0, 12, design_z=11, delta_z=1, has_design_z=True),
        SurveyPoint("B", 1, 0, 8, design_z=9, delta_z=-1, has_design_z=True),
        SurveyPoint("C", 0, 1, 10, design_z=10, delta_z=0, has_design_z=True),
    ]
    boundary = [(0.0, 0.0), (1.0, 0.0), (1.0, 1.0), (0.0, 1.0)]
    path = tmp_path / "job.tinproj.json"

    save_project(
        path,
        points,
        boundary,
        design_elevation=10,
        use_partition=True,
        partition={"A": 11, "B": 9, "C": 10},
        project_name="现场试验段",
    )
    loaded = load_project(path)

    assert loaded["project_name"] == "现场试验段"
    assert loaded["boundary"] == boundary
    assert loaded["use_partition"] is True
    assert loaded["partition"]["A"] == 11
    by_id = {point.id: point for point in loaded["points"]}
    assert by_id["B"].design_z == 9
    assert by_id["B"].has_design_z


def test_export_dxf_writes_boundary_and_zero_contour(tmp_path):
    points = [
        SurveyPoint("A", 0, 0, 1),
        SurveyPoint("B", 1, 0, -1),
        SurveyPoint("C", 0, 1, -1),
    ]
    calculator = TINEarthworkCalculator(0)
    calculator.add_points(points)
    calculator.set_boundary([(0, 0), (1, 0), (0, 1)])
    result = calculator.run_full_calculation()
    path = tmp_path / "earthwork.dxf"

    assert DataExporter().export_boundary_dxf(result, path)
    text = path.read_text(encoding="utf-8")
    assert "$ACADVER\n1\nAC1009" in text  # R12，AutoCAD/CASS 可直接打开
    assert "LWPOLYLINE" not in text
    assert text.count("\nPOLYLINE\n") == 2
    assert "BOUNDARY" in text
    assert "ZERO_CONTOUR" in text
    assert "0.5" in text


def test_import_boundary_from_dxf_lwpolyline_and_polyline(tmp_path):
    lwpolyline = tmp_path / "lw.dxf"
    lwpolyline.write_text(
        "0\nSECTION\n2\nENTITIES\n"
        "0\nLWPOLYLINE\n8\nROAD\n90\n3\n70\n1\n10\n0\n20\n0\n10\n1\n20\n0\n10\n1\n20\n1\n"
        "0\nLWPOLYLINE\n8\nBOUNDARY\n90\n4\n70\n1\n"
        "10\n0\n20\n0\n10\n10\n20\n0\n10\n10\n20\n10\n10\n0\n20\n10\n"
        "0\nENDSEC\n0\nEOF\n",
        encoding="utf-8",
    )
    assert import_boundary_from_dxf(lwpolyline) == [(0.0, 0.0), (10.0, 0.0), (10.0, 10.0), (0.0, 10.0)]

    polyline = tmp_path / "pl.dxf"
    polyline.write_text(
        "0\nSECTION\n2\nENTITIES\n"
        "0\nPOLYLINE\n8\n红线\n66\n1\n70\n1\n10\n0\n20\n0\n"
        "0\nVERTEX\n8\n红线\n10\n2\n20\n2\n"
        "0\nVERTEX\n8\n红线\n10\n8\n20\n2\n"
        "0\nVERTEX\n8\n红线\n10\n8\n20\n8\n"
        "0\nVERTEX\n8\n红线\n10\n2\n20\n8\n"
        "0\nSEQEND\n"
        "0\nENDSEC\n0\nEOF\n",
        encoding="utf-8",
    )
    assert import_boundary_from_dxf(polyline) == [(2.0, 2.0), (8.0, 2.0), (8.0, 8.0), (2.0, 8.0)]

    exported = tmp_path / "roundtrip.dxf"
    result, _ = _result()
    result.boundary_points = [(0.0, 0.0), (5.0, 0.0), (5.0, 4.0), (0.0, 4.0)]
    assert DataExporter().export_boundary_dxf(result, exported)
    assert import_boundary_from_dxf(exported) == result.boundary_points


def test_import_cass_dat_and_delimited_txt(tmp_path):
    dat = tmp_path / "site.dat"
    dat.write_text("1,DMD,100.5,200.5,12.0\n2,,101.5,201.5,11.0\n", encoding="gbk")
    points, issues, _ = DataImporter.import_file(dat)
    by_id = {point.id: point for point in points}
    assert by_id["1"].x == 100.5
    assert by_id["1"].y == 200.5
    assert by_id["1"].z == 12.0
    assert by_id["2"].x == 101.5
    assert not issues["format_errors"]

    space_txt = tmp_path / "points.txt"
    space_txt.write_text("A 0 0 12\nB 1 0 8\nC 0 1 10\n", encoding="utf-8")
    space_points, _, _ = DataImporter.import_file(space_txt)
    assert [(p.id, p.x, p.y, p.z) for p in space_points] == [
        ("A", 0.0, 0.0, 12.0),
        ("B", 1.0, 0.0, 8.0),
        ("C", 0.0, 1.0, 10.0),
    ]

    comma_txt = tmp_path / "comma.txt"
    comma_txt.write_text("点号,X,Y,Z\nA,0,0,12\nB,1,0,8\n", encoding="utf-8")
    comma_points, _, _ = DataImporter.import_file(comma_txt)
    assert comma_points[0].id == "A"
    assert comma_points[1].x == 1.0


def test_export_pdf_report_has_summary_and_figure(tmp_path):
    result, points = _result()
    result.boundary_points = [(0.0, 0.0), (1.0, 0.0), (0.0, 1.0)]
    path = tmp_path / "report.pdf"
    assert DataExporter("试验段").export_pdf_report(
        result, len(points), points, result.boundary_points, path
    )
    data = path.read_bytes()
    assert data.startswith(b"%PDF")
    assert path.stat().st_size > 1000
    assert data.count(b"/Type /Page") >= 2 or b"/Count 2" in data


def test_exports_are_created_and_readable(tmp_path):
    result, points = _result()
    exporter = DataExporter("测试项目")
    excel_path = tmp_path / "report.xlsx"
    csv_path = tmp_path / "detail.csv"

    assert exporter.export_summary_excel(result, excel_path, points)
    assert exporter.export_triangles_csv(result, csv_path)
    workbook = openpyxl.load_workbook(excel_path)
    assert workbook.sheetnames == ["土方量汇总表", "三角形计算明细", "异常数据检查"]
    assert workbook["土方量汇总表"]["A1"].value == "测试项目"
    assert not pd.read_csv(csv_path, encoding="utf-8-sig").empty
    assert "总挖方量" in exporter.generate_report_text(result, len(points))


def _rows(sheet):
    return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]


def test_excel_check_sheet_lists_real_validation_results(tmp_path):
    """回归：原先第 3 页是“见导入日志”的占位文字。"""
    result, points = _result()
    points.append(SurveyPoint("A", 5, 5, 10))  # 重复点号
    path = tmp_path / "report.xlsx"
    DataExporter("试验段").export_summary_excel(result, path, points)
    rows = _rows(openpyxl.load_workbook(path)["异常数据检查"])
    by_name = {row[0]: row for row in rows[1:]}
    assert by_name["重复点号"][1] == 1
    assert "A" in by_name["重复点号"][2]
    assert by_name["缺失值"][2] == "未发现"
    assert "计算提示" in by_name


def test_reports_describe_partition_design_elevation(tmp_path):
    """回归：启用分区高程后，报告不能再写成单一的“设计高程 X m”。"""
    points = [
        SurveyPoint("A", 0, 0, 12),
        SurveyPoint("B", 10, 0, 8),
        SurveyPoint("C", 0, 10, 10),
    ]
    calculator = TINEarthworkCalculator()
    calculator.add_points(points)
    calculator.set_design_elevations({"A": 11}, default=9.5)
    result = calculator.run_full_calculation()
    exporter = DataExporter("分区试验")
    text = exporter.generate_report_text(result, len(points))
    assert "9.500 ~ 11.000" in text
    path = tmp_path / "partition.xlsx"
    exporter.export_summary_excel(result, path, points)
    summary = {row[0]: row[1] for row in _rows(openpyxl.load_workbook(path)["土方量汇总表"]) if row}
    assert "9.500 ~ 11.000" in summary["设计高程"]


def test_report_includes_coverage_and_warnings():
    points = [SurveyPoint(f"P{i}", x, y, 11.0) for i, (x, y) in enumerate([(0, 0), (100, 0), (0, 100), (100, 100)])]
    calculator = TINEarthworkCalculator(10)
    calculator.add_points(points)
    calculator.set_boundary([(-50, -50), (150, -50), (150, 150), (-50, 150)])
    result = calculator.run_full_calculation()
    text = DataExporter().generate_report_text(result, len(points))
    assert "测点覆盖率 25.00%" in text
    assert "不在测点覆盖范围内" in text


def test_export_failure_raises_with_reason(tmp_path):
    result, points = _result()
    missing_dir = tmp_path / "no-such-dir" / "report.xlsx"
    try:
        DataExporter().export_summary_excel(result, missing_dir, points)
    except OSError as error:
        assert "no-such-dir" in str(error)
    else:
        raise AssertionError("导出到不存在的目录应当报错")


def test_blank_or_nan_rows_are_rejected_not_imported(tmp_path):
    """回归：空高程曾被解析为 NaN，导致总填方显示 nan。"""
    source = tmp_path / "gaps.xlsx"
    pd.DataFrame({
        "点号": ["A", "B", "C", None],
        "X": [0, 10, 0, 10],
        "Y": [0, 0, 10, 10],
        "Z": [11, 12, None, 11.5],
    }).to_excel(source, index=False)
    points, issues, _ = DataImporter.import_file(source)
    assert [point.id for point in points] == ["A", "B", "P4"]
    assert any("第3行" in item for item in issues["format_errors"])

    text = tmp_path / "nan.txt"
    text.write_text("A 0 0 12\nB 1 0 nan\nC 0 1 10\n", encoding="utf-8")
    text_points, text_issues, _ = DataImporter.import_file(text)
    assert [point.id for point in text_points] == ["A", "C"]
    assert text_issues["format_errors"]


def test_three_column_file_maps_xyz_without_id(tmp_path):
    source = tmp_path / "xyz.csv"
    source.write_text("东,北,高\n0,0,10\n1,0,11\n0,1,12\n", encoding="utf-8")
    points, issues, _ = DataImporter.import_file(source)
    assert [(p.id, p.x, p.y, p.z) for p in points] == [
        ("P1", 0.0, 0.0, 10.0), ("P2", 1.0, 0.0, 11.0), ("P3", 0.0, 1.0, 12.0)
    ]
    assert not issues["format_errors"]


def test_project_file_is_written_atomically(tmp_path, monkeypatch):
    path = tmp_path / "job.tinproj.json"
    save_project(path, [SurveyPoint("A", 0, 0, 1)], [], project_name="旧")
    original = path.read_text(encoding="utf-8")

    import json as json_module

    def broken_dumps(*args, **kwargs):
        raise RuntimeError("写入中断")

    monkeypatch.setattr(json_module, "dumps", broken_dumps)
    try:
        save_project(path, [SurveyPoint("B", 0, 0, 1)], [], project_name="新")
    except RuntimeError:
        pass
    assert path.read_text(encoding="utf-8") == original
    assert [p.name for p in tmp_path.iterdir()] == ["job.tinproj.json"]
