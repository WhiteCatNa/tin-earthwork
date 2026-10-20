"""不确定度评定：抽稀外推修正内插偏差、95% 区间、补测点推荐、报表与导出。"""
import math

import numpy as np
import pytest
from openpyxl import load_workbook

from core.calculator import COMPARE_DESIGN, SurveyPoint, TINEarthworkCalculator
from core.uncertainty import _area_weights2, _nearest_cv, evaluate_uncertainty
from gui.uncertainty_window import export_recommendations_csv, export_recommendations_dxf
from utils.data_handler import DataExporter, load_project, save_project

SIZE = 100.0
SQUARE = [(0.0, 0.0), (SIZE, 0.0), (SIZE, SIZE), (0.0, SIZE)]


def hill(x, y):
    """场地中间一个高斯山包，设计高程 50 时只有挖方。"""
    return 50 + 3 * np.exp(-((x - 50) ** 2 + (y - 50) ** 2) / (2 * 20 ** 2))


def hill_volume():
    # ∫∫ 3·exp(−r²/2s²) 在 [0,100]² 上：3·2πs²·(erf(50/(s√2)))²
    s = 20.0
    return 3 * 2 * math.pi * s * s * math.erf(50 / (s * math.sqrt(2))) ** 2


def _points(func, n=300, seed=1, noise=0.0, prefix="P"):
    rng = np.random.default_rng(seed)
    xy = np.vstack([SQUARE, rng.uniform(0, SIZE, (n, 2))])
    z = func(xy[:, 0], xy[:, 1]) + rng.normal(0, noise, len(xy)) * (noise > 0)
    return [SurveyPoint(f"{prefix}{i}", float(x), float(y), float(v)) for i, ((x, y), v) in enumerate(zip(xy, z))]


def _calc(func=hill, n=300, seed=1, noise=0.0, holes=None):
    calc = TINEarthworkCalculator(50.0)
    calc.add_points(_points(func, n, seed, noise))
    calc.set_boundary(SQUARE)
    if holes:
        calc.set_holes(holes)
    return calc, calc.run_full_calculation()


def test_bias_correction_moves_toward_truth():
    truth = hill_volume()
    errors = []
    for seed in range(4):
        calc, result = _calc(n=200, seed=seed)
        u = evaluate_uncertainty(calc, result, sigma_measure=0.0, seed=seed)
        assert u.tin[0] == pytest.approx(result.total_cut)
        assert u.bias[0] < 0                 # 山包被削平：TIN 法挖方偏小
        assert abs(u.estimate[0] - truth) < 1.96 * u.sigma[0]
        errors.append((abs(result.total_cut - truth), abs(u.estimate[0] - truth)))
    before, after = np.array(errors).mean(axis=0)
    assert after < 0.8 * before


def test_plane_has_no_interpolation_error():
    calc, result = _calc(lambda x, y: 49 + 0.02 * x + 0.01 * y)
    u = evaluate_uncertainty(calc, result, sigma_measure=0.01, recommend=5)
    assert all(abs(b) < 1e-6 for b in u.bias)
    assert all(s < 1e-6 for s in u.interp_sigma)
    # 只有测量误差：σ ≈ σm·√Σw²，约为 σm·A/√n
    assert u.sigma[2] == pytest.approx(u.measure_sigma[2])
    assert 0.3 < u.sigma[2] / (0.01 * SIZE * SIZE / math.sqrt(304)) < 3
    assert not u.recommendations          # 留一残差只是测量误差，不推荐
    assert any("看不出内插误差" in w for w in u.warnings)


def test_recommendations_inside_and_decreasing():
    hole = [(60.0, 60.0), (80.0, 60.0), (80.0, 80.0), (60.0, 80.0)]
    calc, result = _calc(n=250, holes=[hole])
    u = evaluate_uncertainty(calc, result, sigma_measure=0.005, recommend=15)
    assert len(u.recommendations) == 15
    remaining = [r.remaining for r in u.recommendations]
    assert remaining == sorted(remaining, reverse=True) and 0 < remaining[-1] < 1
    gains = [r.reduction for r in u.recommendations]
    assert all(g > 0 for g in gains)
    for r in u.recommendations:
        assert 0 <= r.x <= SIZE and 0 <= r.y <= SIZE
        assert not (60 < r.x < 80 and 60 < r.y < 80)
    assert u.triangle_sigma and u.indicator > 0


def test_recommended_points_beat_random_points():
    truth = hill_volume()
    gains = []
    for seed in range(3):
        rng = np.random.default_rng(100 + seed)
        points = _points(hill, 150, seed)
        calc = TINEarthworkCalculator(50.0)
        calc.add_points(points)
        calc.set_boundary(SQUARE)
        result = calc.run_full_calculation()
        u = evaluate_uncertainty(calc, result, sigma_measure=0.0, recommend=20, seed=seed)
        errors = []
        for extra in ([(r.x, r.y) for r in u.recommendations], rng.uniform(0, SIZE, (20, 2))):
            more = points + [SurveyPoint(f"N{k}", float(x), float(y), float(hill(x, y))) for k, (x, y) in enumerate(extra)]
            again = TINEarthworkCalculator(50.0)
            again.add_points(more)
            again.set_boundary(SQUARE)
            errors.append(abs(again.run_full_calculation().total_cut - truth))
        gains.append(errors)
    recommended, random_points = np.array(gains).mean(axis=0)
    assert recommended < random_points


def test_too_few_points_measurement_only():
    calc, result = _calc(n=8)
    u = evaluate_uncertainty(calc, result, sigma_measure=0.02)
    assert u.repeats == 0 and u.bias == (0.0, 0.0, 0.0)
    assert u.sigma[2] > 0
    assert any("内部测点少于" in w for w in u.warnings)


def test_compare_modes():
    calc = TINEarthworkCalculator()
    calc.add_points(_points(hill, 200, 1))
    calc.set_compare_points(_points(lambda x, y: 50 + 0 * x, 150, 2, prefix="Q"))
    calc.set_boundary(SQUARE)
    result = calc.run_full_calculation()
    u = evaluate_uncertainty(calc, result, sigma_measure=0.01, recommend=3)
    assert u.repeats >= 2 and u.bias[0] < 0
    assert len(u.recommendations) == 0 or len(u.recommendations) == 3   # 后期为平面，不一定推荐
    # 两期对比测量误差计两次，大于只有一期
    single, _ = _calc(n=200)
    assert u.measure_sigma[2] > evaluate_uncertainty(single, single.run_full_calculation(), 0.01).measure_sigma[2]

    design = TINEarthworkCalculator()
    design.add_points(_points(hill, 200, 1))
    design.set_compare_points(_points(lambda x, y: 50 + 0 * x, 30, 3, prefix="D"), kind=COMPARE_DESIGN)
    design.set_boundary(SQUARE)
    result = design.run_full_calculation()
    u = evaluate_uncertainty(design, result, sigma_measure=0.01, recommend=3)
    assert u.bias[0] < 0 and len(u.recommendations) == 3


def test_helpers_and_validation():
    calc, result = _calc(n=50)
    with pytest.raises(ValueError):
        evaluate_uncertainty(calc, result, sigma_measure=-1)
    with pytest.raises(ValueError):
        evaluate_uncertainty(calc, result, recommend=5000)
    grid = [SurveyPoint(f"G{i}", float(i % 10), float(i // 10), 0.0) for i in range(100)]
    assert _nearest_cv(grid) < 0.05
    assert _nearest_cv(_points(hill, 300)) > 0.3
    # 1 m 格网：64 个内部点各分担 1 m²，边上的点分担更少，Σw = 81
    assert 64 < _area_weights2(grid, 81.0) < 81
    assert _area_weights2(grid, 162.0) == pytest.approx(4 * _area_weights2(grid, 81.0))


def test_reports_and_exports(tmp_path):
    calc, result = _calc(n=200)
    result.uncertainty = evaluate_uncertainty(calc, result, sigma_measure=0.01, recommend=4)
    rows = DataExporter.summary_rows(result)
    assert any(row[0] == "不确定度评定" for row in rows)
    path = tmp_path / "u.xlsx"
    assert DataExporter("测试").export_summary_excel(result, str(path), calc.points)
    sheet = load_workbook(path)["不确定度评定"]
    values = [c.value for row in sheet.iter_rows() for c in row if c.value is not None]
    assert "挖方" in values and "R4" in values
    assert "【不确定度评定】" in DataExporter("测试").generate_report_text(result, len(calc.points))

    csv_path = tmp_path / "r.csv"
    export_recommendations_csv(result.uncertainty, str(csv_path))
    assert csv_path.read_text(encoding="utf-8-sig").splitlines()[1].startswith("R1,")
    dxf = tmp_path / "r.dxf"
    export_recommendations_dxf(result.uncertainty, str(dxf))
    assert dxf.read_text().count("POINT") == 4

    project = tmp_path / "p.tinproj.json"
    save_project(str(project), calc.points, SQUARE, uncertainty={"enabled": True, "sigma_measure": 0.01, "recommend": 4})
    assert load_project(str(project))["uncertainty"]["recommend"] == 4
