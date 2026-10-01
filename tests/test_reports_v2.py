"""2.0 新增输出：两期对比用语、方格网校核报告与 DXF、工程文件 v2（兼容 v1）。"""
import json

from openpyxl import load_workbook

from core.calculator import SurveyPoint, TINEarthworkCalculator
from core.grid_check import run_grid_check
from core.surface import PlaneDesign
from utils.data_handler import DataExporter, load_project, save_project
from utils.dxf_io import import_boundary_from_dxf

BOUNDARY = [(0.0, 0.0), (20.0, 0.0), (20.0, 20.0), (0.0, 20.0)]


def _square(z, prefix="P", size=20.0):
    return [SurveyPoint(f"{prefix}{i}", x, y, z) for i, (x, y) in
            enumerate([(0, 0), (size, 0), (size, size), (0, size), (size / 2, size / 3)], 1)]


def _compare_result(grid=True):
    calculator = TINEarthworkCalculator()
    calculator.add_points(_square(10.0))
    calculator.set_compare_points(_square(9.0, prefix="Q"))
    calculator.set_boundary(BOUNDARY)
    result = calculator.run_full_calculation()
    if grid:
        result.grid_check = run_grid_check(calculator, result, 5.0)
    return result


def test_compare_mode_uses_before_after_wording(tmp_path):
    result = _compare_result(grid=False)
    rows = {row[0]: row for row in DataExporter.summary_rows(result)}
    assert "前期高于后期" in rows["总挖方量"][3]
    assert rows["比较面"][1].startswith("两期对比")
    assert rows["后期测点数"][1] == 5
    headers = DataExporter.detail_headers(result)
    assert "顶点1前期高程" in headers and "顶点1后期高程" in headers
    text = DataExporter("两期").generate_report_text(result, 5)
    assert "前期测点数: 5" in text and "后期测点数: 5" in text and "两期TIN叠加" in text


def test_grid_check_in_summary_excel_text_and_dxf(tmp_path):
    result = _compare_result()
    rows = {row[0]: row for row in DataExporter.summary_rows(result)}
    assert rows["方格网法挖方"][1] == 400.0 and "+0.00%" in rows["方格网法挖方"][3]

    exporter = DataExporter("方格网")
    excel = tmp_path / "report.xlsx"
    exporter.export_summary_excel(result, excel, _square(10.0))
    workbook = load_workbook(excel)
    assert workbook.sheetnames[-3:] == ["方格网校核", "方格网角点", "方格网方格"]
    corner_rows = list(workbook["方格网角点"].iter_rows(values_only=True))
    assert corner_rows[0][5:8] == ("前期高程", "后期高程", "施工高度(m)")
    assert len(corner_rows) == 1 + 25
    assert len(list(workbook["方格网方格"].iter_rows(values_only=True))) == 1 + 16

    text = exporter.generate_report_text(result, 5)
    assert "【方格网校核】" in text and "方格边长: 5 m" in text

    dxf = tmp_path / "grid.dxf"
    exporter.export_boundary_dxf(result, dxf)
    content = dxf.read_text(encoding="utf-8")
    assert content.count("\nTEXT\n8\nGRID_") == 25 * 2 + 16
    assert content.count("\nTEXT\n8\nBOUNDARY_POINT\n") == len(result.boundary_points)
    for layer in ("GRID", "GRID_HEIGHT", "GRID_ELEV", "GRID_VOLUME"):
        assert f"\n8\n{layer}\n" in content
    # 方格线不会被当成计算边界读回
    assert import_boundary_from_dxf(dxf) == BOUNDARY


def test_pdf_report_with_grid_check(tmp_path):
    path = tmp_path / "book.pdf"
    DataExporter("方格网").export_pdf_report(_compare_result(), 5, _square(10.0), BOUNDARY, path)
    assert path.read_bytes().startswith(b"%PDF")


def test_project_v2_roundtrip(tmp_path):
    path = tmp_path / "site.tinproj.json"
    plane = PlaneDesign(1.0, 2.0, 3.0, 0.5, -0.3)
    save_project(
        path, _square(10.0), BOUNDARY, design_elevation=9.5, project_name="二期",
        design_mode="plane", design_plane=plane.to_dict(), use_partition=True, partition={"P1": 9.0},
        compare_points=_square(9.0, prefix="Q"), compare_source="after.xlsx",
        grid_enabled=True, grid_spacing=5.0,
    )
    data = load_project(path)
    assert data["design_mode"] == "plane"
    assert PlaneDesign.from_dict(data["design_plane"]) == plane
    assert data["partition"] == {"P1": 9.0} and data["use_partition"]
    assert [p.id for p in data["compare_points"]] == ["Q1", "Q2", "Q3", "Q4", "Q5"]
    assert data["compare_source"] == "after.xlsx"
    assert (data["grid_enabled"], data["grid_spacing"]) == (True, 5.0)


def test_project_v1_file_still_opens(tmp_path):
    path = tmp_path / "old.tinproj.json"
    path.write_text(json.dumps({
        "format": "tin-earthwork-project", "format_version": 1, "app_version": "1.3.0",
        "project_name": "旧工程", "design_elevation": 12.5, "use_partition": False, "partition": {},
        "boundary": [[0, 0], [1, 0], [1, 1]],
        "points": [{"id": "A", "x": 0, "y": 0, "z": 13}],
    }, ensure_ascii=False), encoding="utf-8")
    data = load_project(path)
    assert data["design_mode"] == "flat" and data["design_elevation"] == 12.5
    assert data["compare_points"] == [] and data["design_plane"] is None
    assert data["grid_enabled"] is False


def test_project_from_newer_version_is_rejected(tmp_path):
    path = tmp_path / "future.tinproj.json"
    path.write_text(json.dumps({"format": "tin-earthwork-project", "format_version": 99, "points": []}),
                    encoding="utf-8")
    try:
        load_project(path)
    except ValueError as error:
        assert "更新版本" in str(error)
    else:
        raise AssertionError("应拒绝更新版本的工程文件")
