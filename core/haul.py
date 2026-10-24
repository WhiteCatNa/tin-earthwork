"""
土方调配：把计算范围划成方格调配区，统计每区挖方、填方及其重心，
区内先自行平衡，余缺部分按“运量 × 运距”最小求调配方案（运输问题，线性规划求解）。

每区的挖方、填方由三角形按方格精确切分后积分得到，重心是体积加权的精确重心；
运距取挖方区挖方重心到填方区填方重心的直线距离。考虑可松性时，填方区需要的自然方 = 填方 ÷ Ks'。
挖方多于填方需要时，余土外运（弃土）；不足时由外部借土，二者都不计入场内运距。
"""
import math
import string
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import numpy as np

from core.calculator import TINEarthworkCalculator
from core.geometry import clip_ring_halfplane
from core.layers import surface_pieces

XY = Tuple[float, float]
MAX_ZONES = 900
MIN_MOVE = 1e-6           # 小于该方量的调配不列出 (m³)


@dataclass
class HaulZone:
    name: str
    row: int
    col: int
    bounds: Tuple[float, float, float, float]      # xmin, ymin, xmax, ymax（全局坐标）
    cut: float = 0.0
    fill: float = 0.0
    cut_center: Optional[XY] = None
    fill_center: Optional[XY] = None
    local: float = 0.0          # 区内自行平衡的自然方
    supply: float = 0.0         # 可调出的自然方
    demand: float = 0.0         # 需调入的自然方


@dataclass
class HaulMove:
    source: str
    target: str                 # 目标区名称；“弃土”表示外运
    volume: float               # 自然方 (m³)
    distance: float             # 运距 (m)；弃土、借土为 0（不计场内运距）

    @property
    def work(self) -> float:
        return self.volume * self.distance


@dataclass
class HaulResult:
    zone_size: float
    compact: float
    zones: List[HaulZone] = field(default_factory=list)
    moves: List[HaulMove] = field(default_factory=list)
    waste: float = 0.0          # 余土外运（自然方）
    borrow: float = 0.0         # 外部借土（自然方）

    @property
    def hauled(self) -> float:
        return sum(m.volume for m in self.moves if m.target != WASTE and m.source != BORROW)

    @property
    def work(self) -> float:
        """场内总运输量 (m³·m)。"""
        return sum(m.work for m in self.moves)

    @property
    def average_distance(self) -> float:
        hauled = self.hauled
        return self.work / hauled if hauled > 0 else 0.0

    def zone(self, name: str) -> HaulZone:
        return next(z for z in self.zones if z.name == name)


WASTE = "弃土"
BORROW = "借土"


def _zone_name(row: int, col: int) -> str:
    """列用字母（A、B…、AA…），行用数字，与图上一致：A1 在左下角。"""
    letters = ""
    col += 1
    while col:
        col, rest = divmod(col - 1, 26)
        letters = string.ascii_uppercase[rest] + letters
    return f"{letters}{row + 1}"


def _moments(polygon: np.ndarray, plane: Tuple[float, float, float]) -> Tuple[float, float, float]:
    """∫h、∫x·h、∫y·h（h 为线性函数，多边形按扇形剖分，有向面积保证凹多边形精确）。"""
    if len(polygon) < 3:
        return 0.0, 0.0, 0.0
    a, b, c = plane
    h = polygon[:, 0] * a + polygon[:, 1] * b + c
    p0, h0 = polygon[0], h[0]
    total = mx = my = 0.0
    for i in range(1, len(polygon) - 1):
        p1, p2 = polygon[i], polygon[i + 1]
        area = ((p1[0] - p0[0]) * (p2[1] - p0[1]) - (p2[0] - p0[0]) * (p1[1] - p0[1])) / 2
        hs = (h0, h[i], h[i + 1])
        xs = (p0[0], p1[0], p2[0])
        ys = (p0[1], p1[1], p2[1])
        sum_h = sum(hs)
        total += area * sum_h / 3
        # ∫_T f·g = A/12 · (Σ f_i g_i + Σf_i · Σg_i)
        mx += area / 12 * (sum(x * v for x, v in zip(xs, hs)) + sum(xs) * sum_h)
        my += area / 12 * (sum(y * v for y, v in zip(ys, hs)) + sum(ys) * sum_h)
    if total < 0:
        return -total, -mx, -my
    return total, mx, my


def suggest_zone_size(calculator: TINEarthworkCalculator, target: int = 36) -> float:
    """调配区边长：使计算范围大约分成 target 个区，取 5/10/20/25/50×10ⁿ 的整齐值。"""
    polygons = [p for p, _, _, sign in surface_pieces(calculator) if sign > 0]
    if not polygons:
        return 10.0
    xy = np.concatenate(polygons)
    span = np.ptp(xy, axis=0)
    raw = math.sqrt(max(float(span[0] * span[1]), 1e-6) / target)
    exponent = math.floor(math.log10(raw))
    for nice in (1, 2, 2.5, 5, 10):
        if nice * 10 ** exponent >= raw:
            return float(nice * 10 ** exponent)
    return float(10 ** (exponent + 1))


def run_haul(calculator: TINEarthworkCalculator, zone_size: float, compact: float = 1.0) -> HaulResult:
    """按方格调配区求最小运输量的调配方案（calculator 须已完成计算）。"""
    zone_size = float(zone_size)
    compact = float(compact)
    if not math.isfinite(zone_size) or zone_size <= 0:
        raise ValueError("调配区边长必须大于 0")
    if not math.isfinite(compact) or compact <= 0:
        raise ValueError("最终可松性系数 Ks' 必须大于 0")
    pieces = list(surface_pieces(calculator))
    if not pieces:
        raise ValueError("计算范围内没有三角形，无法调配")
    ox, oy = calculator.origin
    xy = np.concatenate([p for p, _, _, sign in pieces if sign > 0])
    xmin, ymin = xy.min(axis=0)
    xmax, ymax = xy.max(axis=0)
    cols = max(1, math.ceil((xmax - xmin) / zone_size - 1e-9))
    rows = max(1, math.ceil((ymax - ymin) / zone_size - 1e-9))
    if rows * cols > MAX_ZONES:
        raise ValueError(f"调配区边长 {zone_size:g} m 过小：需要 {rows * cols} 个区，请加大边长")

    # [行, 列, (挖 ∫h, ∫xh, ∫yh, 填 ∫h, ∫xh, ∫yh)]，局部坐标
    acc = np.zeros((rows, cols, 6))
    for polygon, natural, design, sign in pieces:
        diff = tuple(p - q for p, q in zip(natural, design))
        neg = tuple(-v for v in diff)
        c0 = max(0, int(math.floor((polygon[:, 0].min() - xmin) / zone_size)))
        c1 = min(cols - 1, int(math.floor((polygon[:, 0].max() - xmin) / zone_size)))
        r0 = max(0, int(math.floor((polygon[:, 1].min() - ymin) / zone_size)))
        r1 = min(rows - 1, int(math.floor((polygon[:, 1].max() - ymin) / zone_size)))
        for r in range(r0, r1 + 1):
            for c in range(c0, c1 + 1):
                x0, y0 = xmin + c * zone_size, ymin + r * zone_size
                part = polygon
                # 首行、首列及末行、末列不裁外侧：外侧本来就没有计算范围（避免舍入丢掉边上的碎片）
                if c > 0:
                    part = clip_ring_halfplane(part, 1.0, 0.0, -x0)
                if c < cols - 1:
                    part = clip_ring_halfplane(part, -1.0, 0.0, x0 + zone_size)
                if r > 0:
                    part = clip_ring_halfplane(part, 0.0, 1.0, -y0)
                if r < rows - 1:
                    part = clip_ring_halfplane(part, 0.0, -1.0, y0 + zone_size)
                if len(part) < 3:
                    continue
                for offset, plane in ((0, diff), (3, neg)):
                    clipped = clip_ring_halfplane(part, *plane)
                    if len(clipped) >= 3:
                        acc[r, c, offset:offset + 3] += sign * np.array(_moments(clipped, plane))

    result = HaulResult(zone_size=zone_size, compact=compact)
    for r in range(rows):
        for c in range(cols):
            cut, cmx, cmy, fill, fmx, fmy = acc[r, c]
            cut, fill = max(cut, 0.0), max(fill, 0.0)
            if cut <= MIN_MOVE and fill <= MIN_MOVE:
                continue
            x0, y0 = xmin + c * zone_size + ox, ymin + r * zone_size + oy
            zone = HaulZone(_zone_name(r, c), r, c, (x0, y0, x0 + zone_size, y0 + zone_size), cut, fill)
            if cut > MIN_MOVE:
                zone.cut_center = (cmx / cut + ox, cmy / cut + oy)
            if fill > MIN_MOVE:
                zone.fill_center = (fmx / fill + ox, fmy / fill + oy)
            need = fill / compact
            zone.local = min(cut, need)
            zone.supply = cut - zone.local
            zone.demand = need - zone.local
            result.zones.append(zone)

    sources = [z for z in result.zones if z.supply > MIN_MOVE]
    sinks = [z for z in result.zones if z.demand > MIN_MOVE]
    supply = np.array([z.supply for z in sources])
    demand = np.array([z.demand for z in sinks])
    moved = np.zeros((len(sources), len(sinks)))
    if len(sources) and len(sinks):
        distance = np.array([[math.dist(s.cut_center, t.fill_center) for t in sinks] for s in sources])
        moved = _transport(supply, demand, distance)
        for i, s in enumerate(sources):
            for j, t in enumerate(sinks):
                if moved[i, j] > MIN_MOVE:
                    result.moves.append(HaulMove(s.name, t.name, float(moved[i, j]), float(distance[i, j])))
    for i, s in enumerate(sources):
        left = float(supply[i] - moved[i].sum()) if len(sinks) else float(supply[i])
        if left > MIN_MOVE:
            result.moves.append(HaulMove(s.name, WASTE, left, 0.0))
            result.waste += left
    for j, t in enumerate(sinks):
        short = float(demand[j] - moved[:, j].sum()) if len(sources) else float(demand[j])
        if short > MIN_MOVE:
            result.moves.append(HaulMove(BORROW, t.name, short, 0.0))
            result.borrow += short
    return result


def _transport(supply: np.ndarray, demand: np.ndarray, distance: np.ndarray) -> np.ndarray:
    """运输问题：场内调运量取最大（= min(总供, 总需)），在此前提下运量×运距最小。"""
    from scipy.optimize import linprog

    m, n = distance.shape
    total = min(supply.sum(), demand.sum())
    a_ub = np.zeros((m + n, m * n))
    for i in range(m):
        a_ub[i, i * n:(i + 1) * n] = 1           # 调出 ≤ 可供
    for j in range(n):
        a_ub[m + j, j::n] = 1                    # 调入 ≤ 需要
    b_ub = np.concatenate([supply, demand])
    a_eq = np.ones((1, m * n))                   # 总调运量 = min(总供, 总需)
    scale = max(float(distance.max()), 1.0)
    solution = linprog(distance.ravel() / scale, A_ub=a_ub, b_ub=b_ub, A_eq=a_eq, b_eq=[total],
                       bounds=(0, None), method="highs")
    if not solution.success:
        raise ValueError(f"调配方案求解失败：{solution.message}")
    return np.clip(solution.x.reshape(m, n), 0, None)


def haul_rows(result: HaulResult) -> Tuple[List[list], List[list]]:
    """(调配区表, 调配方案表)，用于界面和 Excel。"""
    zones = [[z.name, round(z.cut, 3), round(z.fill, 3), round(z.local, 3), round(z.supply, 3), round(z.demand, 3),
              "" if z.cut_center is None else f"{z.cut_center[0]:.2f}, {z.cut_center[1]:.2f}",
              "" if z.fill_center is None else f"{z.fill_center[0]:.2f}, {z.fill_center[1]:.2f}"]
             for z in result.zones]
    moves = [[m.source, m.target, round(m.volume, 3), round(m.distance, 2) if m.distance else "",
              round(m.work, 1) if m.distance else ""] for m in result.moves]
    return zones, moves


ZONE_HEADERS = ["调配区", "挖方(m³)", "填方(m³)", "区内平衡(自然方)", "可调出(自然方)", "需调入(自然方)",
                "挖方重心 X, Y", "填方重心 X, Y"]
MOVE_HEADERS = ["从", "到", "调运量(自然方 m³)", "运距(m)", "运输量(m³·m)"]


def export_haul_excel(result: HaulResult, filepath: str, project_name: str = "") -> None:
    import io
    from pathlib import Path

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "土方调配"
    ws.append([f"{project_name} 土方调配".strip()])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"调配区边长 {result.zone_size:g} m，最终可松性系数 Ks' = {result.compact:g}；"
               f"场内调运 {result.hauled:.2f} m³，总运输量 {result.work:.1f} m³·m，平均运距 {result.average_distance:.2f} m；"
               f"余土外运 {result.waste:.2f} m³，外部借土 {result.borrow:.2f} m³（均为自然方）"])
    ws.append([])
    zones, moves = haul_rows(result)
    ws.append(MOVE_HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in moves:
        ws.append(row)
    ws.append([])
    ws.append(ZONE_HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in zones:
        ws.append(row)
    for column, width in zip("ABCDEFGH", (10, 10, 18, 16, 16, 16, 24, 24)):
        ws.column_dimensions[column].width = width
    buffer = io.BytesIO()
    wb.save(buffer)
    Path(filepath).write_bytes(buffer.getvalue())


def zone_grid(result: HaulResult) -> List[Tuple[str, Tuple[float, float, float, float]]]:
    return [(z.name, z.bounds) for z in result.zones]


def moves_between_zones(result: HaulResult) -> List[Tuple[XY, XY, float]]:
    """场内调运的箭头：(起点, 终点, 方量)。"""
    out = []
    for move in result.moves:
        if move.target == WASTE or move.source == BORROW:
            continue
        out.append((result.zone(move.source).cut_center, result.zone(move.target).fill_center, move.volume))
    return out


def check_balance(result: HaulResult) -> Dict[str, float]:
    """守恒检查：各区挖方 = 区内平衡 + 调出 + 弃土；填方需要 = 区内平衡 + 调入 + 借土。"""
    out = {}
    for zone in result.zones:
        sent = sum(m.volume for m in result.moves if m.source == zone.name)
        received = sum(m.volume for m in result.moves if m.target == zone.name)
        out[zone.name] = max(abs(zone.cut - zone.local - sent), abs(zone.fill / result.compact - zone.local - received))
    return out


def summary_text(result: HaulResult) -> str:
    return (f"场内调运 {result.hauled:.2f} m³，总运输量 {result.work:.1f} m³·m，平均运距 {result.average_distance:.2f} m；"
            f"余土外运 {result.waste:.2f} m³，外部借土 {result.borrow:.2f} m³")
